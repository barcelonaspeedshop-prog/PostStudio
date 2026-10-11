"""md_docx: turn a script / shot list / AI prompts .md into a .docx Mike can open and edit in Pages or Word.

    ~/.venvs/poststudio/bin/python studio/md_docx.py docs/videos/06-istanbul/script.md "/Volumes/PostDrive/.../Istanbul script.docx"

Handles what these files use: # / ## / ### headings, > quotes (the AI prompts), - bullets, **bold**, plain paragraphs.
"""
import re, sys
from docx import Document
from docx.shared import Pt

def runs(p, text):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part: p.add_run(part.replace("`", "")).bold = bool(i % 2)

def convert(src, dst):
    doc = Document(); doc.styles["Normal"].font.size = Pt(12)
    for line in open(src, encoding="utf-8").read().splitlines():
        s = line.strip()
        if not s or s == "---": continue
        if m := re.match(r"(#{1,3})\s+(.*)", s): doc.add_heading(m.group(2).replace("`", ""), level=len(m.group(1)) - 1 or 0); continue
        if s.startswith(">"): p = doc.add_paragraph(style="Intense Quote"); runs(p, s.lstrip("> ")); continue
        if s.startswith("- "): runs(doc.add_paragraph(style="List Bullet"), s[2:]); continue
        runs(doc.add_paragraph(), s)
    doc.save(dst)

if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2]); print(sys.argv[2])
